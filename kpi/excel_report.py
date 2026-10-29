"""Write the biweekly KPI workbook: Summary, Trends (with charts), KPI dictionary, Data checks."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
FILL = {"Green": "D4EDDA", "Amber": "FFF3CD", "Red": "F8D7DA", "—": "F2F2F2"}
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
FMT = {"int": "#,##0", "pct": "0.0%", "dec1": "0.0"}


def _header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font = PatternFill("solid", fgColor=NAVY), Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _arrow(row) -> str:
    if row["change"] is None or pd.isna(row["change"]) or row["change"] == 0:
        return "→"
    up = row["change"] > 0
    good = up == (row["direction"] == "higher_is_better")
    return ("▲" if up else "▼") + (" good" if good else " worse")


def write(path: Path, period, kpis: pd.DataFrame, trend: pd.DataFrame, checks: pd.DataFrame, cfg: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Biweekly Performance Report"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Period: {period.label}   |   compared with {period.previous().label}"
    passed = (checks.status == "PASS").all()
    ws["A3"] = f"Data checks: {'all passed' if passed else 'FAILED - see Data checks sheet'}"
    ws["A3"].font = Font(bold=True, color="1E7B34" if passed else "B00020")

    r = 5
    for team in ["Support", "Delivery"]:
        ws.cell(row=r, column=1, value=f"{team} team").font = Font(size=13, bold=True, color=NAVY)
        r += 1
        _header(ws, r, ["KPI", "This period", "Last period", "Change", "Target", "Status"])
        r += 1
        for _, k in kpis[kpis.team == team].iterrows():
            vals = [k.kpi, k.current, k.previous, _arrow(k), k.target, k.status]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=None if (isinstance(v, float) and pd.isna(v)) else v)
                c.border = THIN
                if i in (2, 3, 5) and v is not None:
                    c.number_format = FMT[k.format]
                if i == 6:
                    c.fill = PatternFill("solid", fgColor=FILL[k.status])
                    c.alignment = Alignment(horizontal="center")
            r += 1
        r += 1
    for col, w in zip("ABCDEF", [34, 13, 13, 12, 10, 10]):
        ws.column_dimensions[col].width = w
    for sheet_ws in [ws]:
        sheet_ws.page_setup.orientation = "landscape"
        sheet_ws.page_setup.fitToWidth, sheet_ws.page_setup.fitToHeight = 1, 0
        sheet_ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Trends sheet with native Excel line charts
    tr = wb.create_sheet("Trends")
    ids = [k["id"] for k in cfg["kpis"]]
    names = {k["id"]: f'{k["team"]}: {k["name"]}' for k in cfg["kpis"]}
    _header(tr, 1, ["Period ending"] + [names[i] for i in ids])
    for ri, (pe, row) in enumerate(trend.iterrows(), 2):
        tr.cell(row=ri, column=1, value=pe)
        for ci, kid in enumerate(ids, 2):
            c = tr.cell(row=ri, column=ci, value=None if pd.isna(row[kid]) else float(row[kid]))
            c.number_format = FMT[next(k["format"] for k in cfg["kpis"] if k["id"] == kid)]
    tr.column_dimensions["A"].width = 14
    for ci in range(2, len(ids) + 2):
        tr.column_dimensions[get_column_letter(ci)].width = 16
    n = len(trend) + 1
    for i, (kid, anchor) in enumerate([("sla_attainment", "B12"), ("on_time_delivery", "J12"),
                                        ("open_backlog", "B30"), ("median_resolution_hours", "J30")]):
        col = ids.index(kid) + 2
        ch = LineChart()
        ch.title, ch.height, ch.width = names[kid], 7, 14
        ch.add_data(Reference(tr, min_col=col, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(tr, min_col=1, min_row=2, max_row=n))
        ch.legend = None
        for series in ch.series:
            series.smooth = False          # straight segments: don't imply values between periods
        tr.add_chart(ch, anchor)

    # KPI dictionary
    dic = wb.create_sheet("KPI dictionary")
    _header(dic, 1, ["Team", "KPI", "Definition", "Target", "Better when", "Calculation (SQL)"])
    for ri, k in enumerate(cfg["kpis"], 2):
        for ci, v in enumerate([k["team"], k["name"], k["definition"], k.get("target"),
                                "higher" if k["direction"] == "higher_is_better" else "lower", " ".join(k["sql"].split())], 1):
            c = dic.cell(row=ri, column=ci, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", [10, 28, 60, 9, 11, 70]):
        dic.column_dimensions[col].width = w

    # Data checks
    dq = wb.create_sheet("Data checks")
    _header(dq, 1, ["Check", "Status", "Detail"])
    for ri, row in enumerate(checks.itertuples(index=False), 2):
        for ci, v in enumerate(row, 1):
            c = dq.cell(row=ri, column=ci, value=v)
            if ci == 2:
                c.fill = PatternFill("solid", fgColor=FILL["Green" if v == "PASS" else "Red"])
    for col, w in zip("ABC", [44, 10, 40]):
        dq.column_dimensions[col].width = w
    wb.save(path)
